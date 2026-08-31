import io
from datetime import datetime, timedelta
from decimal import Decimal

from django.contrib.admin.views.decorators import staff_member_required
from django.db.models import Sum, Count, Q
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone

from orders.models import Order, OrderItem


def get_date_range(filter_type, start_date=None, end_date=None):
    today = timezone.now().date()
    if filter_type == 'daily':
        return today, today
    elif filter_type == 'weekly':
        start = today - timedelta(days=today.weekday())
        return start, today
    elif filter_type == 'monthly':
        start = today.replace(day=1)
        return start, today
    elif filter_type == 'yearly':
        start = today.replace(month=1, day=1)
        return start, today
    elif filter_type == 'custom' and start_date and end_date:
        return start_date, end_date
    return today, today


@staff_member_required
def sales_report_view(request):
    filter_type = request.GET.get('filter', 'daily')
    start_str = request.GET.get('start_date', '')
    end_str = request.GET.get('end_date', '')

    start_date = None
    end_date = None

    if filter_type == 'custom':
        try:
            start_date = datetime.strptime(start_str, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_str, '%Y-%m-%d').date()
        except ValueError:
            filter_type = 'daily'

    start_date, end_date = get_date_range(filter_type, start_date, end_date)

    orders = Order.objects.filter(
        created_at__date__gte=start_date,
        created_at__date__lte=end_date,
    ).exclude(status='cancelled')

    total_orders = orders.count()
    total_sales = orders.aggregate(total=Sum('total'))['total'] or Decimal('0')
    total_offer_discount = orders.aggregate(
        total=Sum('offer_discount')
    )['total'] or Decimal('0')
    total_coupon_discount = orders.aggregate(
        total=Sum('coupon_discount')
    )['total'] or Decimal('0')
    total_discount = total_offer_discount + total_coupon_discount

    context = {
        'orders': orders.order_by('-created_at'),
        'filter_type': filter_type,
        'start_date': start_date,
        'end_date': end_date,
        'total_orders': total_orders,
        'total_sales': total_sales,
        'total_offer_discount': total_offer_discount,
        'total_coupon_discount': total_coupon_discount,
        'total_discount': total_discount,
        'start_str': start_str,
        'end_str': end_str,
    }
    return render(request, 'reports/sales_report.html', context)


@staff_member_required
def download_pdf_report_view(request):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle,
        Paragraph, Spacer, HRFlowable
    )

    filter_type = request.GET.get('filter', 'daily')
    start_str = request.GET.get('start_date', '')
    end_str = request.GET.get('end_date', '')
    start_date = end_date = None

    if filter_type == 'custom':
        try:
            start_date = datetime.strptime(start_str, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_str, '%Y-%m-%d').date()
        except ValueError:
            filter_type = 'daily'

    start_date, end_date = get_date_range(filter_type, start_date, end_date)

    orders = Order.objects.filter(
        created_at__date__gte=start_date,
        created_at__date__lte=end_date,
    ).exclude(status='cancelled').order_by('-created_at')

    total_sales = orders.aggregate(total=Sum('total'))['total'] or Decimal('0')
    total_discount = (
        (orders.aggregate(o=Sum('offer_discount'))['o'] or Decimal('0')) +
        (orders.aggregate(c=Sum('coupon_discount'))['c'] or Decimal('0'))
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            rightMargin=15*mm, leftMargin=15*mm,
                            topMargin=15*mm, bottomMargin=15*mm)
    story = []

    title_style = ParagraphStyle(
        'Title', fontName='Helvetica-Bold', fontSize=18,
        textColor=colors.HexColor('#4A1836'), spaceAfter=4
    )
    sub_style = ParagraphStyle(
        'Sub', fontName='Helvetica', fontSize=9,
        textColor=colors.HexColor('#C8A96B'), spaceAfter=10
    )
    label_style = ParagraphStyle(
        'Label', fontName='Helvetica-Bold', fontSize=9,
        textColor=colors.HexColor('#4A1836')
    )
    info_style = ParagraphStyle(
        'Info', fontName='Helvetica', fontSize=9,
        textColor=colors.HexColor('#2B2B2B')
    )

    story.append(Paragraph("LIORA — Sales Report", title_style))
    story.append(Paragraph(
        f"Period: {start_date.strftime('%d %b %Y')} to {end_date.strftime('%d %b %Y')}",
        sub_style
    ))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#E8DCCB')))
    story.append(Spacer(1, 6*mm))

    # Summary
    summary_data = [
        [Paragraph("Total Orders", label_style), Paragraph(str(orders.count()), info_style)],
        [Paragraph("Total Sales", label_style), Paragraph(f"₹{total_sales}", info_style)],
        [Paragraph("Total Discounts", label_style), Paragraph(f"₹{total_discount}", info_style)],
    ]
    summary_table = Table(summary_data, colWidths=[60*mm, 60*mm])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#FAF7F2')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E8DCCB')),
        ('PADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 8*mm))

    # Orders table
    header = [
        Paragraph("Order ID", label_style),
        Paragraph("Date", label_style),
        Paragraph("Customer", label_style),
        Paragraph("Discount", label_style),
        Paragraph("Total", label_style),
        Paragraph("Status", label_style),
    ]
    rows = [header]
    for order in orders:
        discount = (order.offer_discount or 0) + (order.coupon_discount or 0)
        rows.append([
            Paragraph(order.order_id, info_style),
            Paragraph(order.created_at.strftime("%d %b %Y"), info_style),
            Paragraph(order.full_name, info_style),
            Paragraph(f"₹{discount}", info_style),
            Paragraph(f"₹{order.total}", info_style),
            Paragraph(order.get_status_display(), info_style),
        ])

    orders_table = Table(rows, colWidths=[35*mm, 22*mm, 40*mm, 22*mm, 22*mm, 22*mm])
    orders_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4A1836')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E8DCCB')),
        ('PADDING', (0, 0), (-1, -1), 5),
        ('ROWBACKGROUND', (0, 1), (-1, -1), colors.white),
    ]))
    story.append(orders_table)

    doc.build(story)
    buffer.seek(0)

    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="LIORA_Sales_Report_{start_date}.pdf"'
    return response


@staff_member_required
def download_excel_report_view(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from decimal import Decimal

    filter_type = request.GET.get('filter', 'daily')
    start_str = request.GET.get('start_date', '')
    end_str = request.GET.get('end_date', '')
    start_date = end_date = None

    if filter_type == 'custom':
        try:
            start_date = datetime.strptime(start_str, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_str, '%Y-%m-%d').date()
        except ValueError:
            filter_type = 'daily'

    start_date, end_date = get_date_range(filter_type, start_date, end_date)

    orders = Order.objects.filter(
        created_at__date__gte=start_date,
        created_at__date__lte=end_date,
    ).exclude(status='cancelled').order_by('-created_at')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Report"

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4A1836")
    center = Alignment(horizontal='center')

    # Title
    ws.merge_cells('A1:F1')
    ws['A1'] = f"LIORA Sales Report — {start_date} to {end_date}"
    ws['A1'].font = Font(bold=True, size=14, color="4A1836")
    ws['A1'].alignment = center

    ws.append([])  # Empty row

    # Summary
    total_sales = orders.aggregate(total=Sum('total'))['total'] or Decimal('0')
    total_discount = (
        (orders.aggregate(o=Sum('offer_discount'))['o'] or Decimal('0')) +
        (orders.aggregate(c=Sum('coupon_discount'))['c'] or Decimal('0'))
    )
    ws.append(["Total Orders", orders.count()])
    ws.append(["Total Sales", float(total_sales)])
    ws.append(["Total Discounts", float(total_discount)])
    ws.append([])

    # Headers
    headers = ["Order ID", "Date", "Customer", "Email",
               "Offer Discount", "Coupon Discount", "Wallet Used",
               "Shipping", "Total", "Payment", "Status"]
    ws.append(headers)

    header_row = ws.max_row
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    # Data rows
    for order in orders:
        ws.append([
            order.order_id,
            order.created_at.strftime("%d %b %Y"),
            order.full_name,
            order.user.email if order.user else '',
            float(order.offer_discount or 0),
            float(order.coupon_discount or 0),
            float(order.wallet_amount_used or 0),
            float(order.shipping_charge or 0),
            float(order.total or 0),
            order.payment_method.upper(),
            order.get_status_display(),
        ])

    # Auto width
    for col in ws.columns:
        max_length = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 4

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="LIORA_Sales_Report_{start_date}.xlsx"'
    return response